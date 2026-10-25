"""Testy rejestru formatów (T8, kontrakt C9)."""

import io
import pickle
import zipfile
from pathlib import Path

import openpyxl
import pytest
from docx import Document
from PIL import Image
from pptx import Presentation
from pptx.util import Emu

from app.parsing import registry
from app.parsing.errors import (
    CorruptedFile,
    EmptyDocument,
    UnsupportedFormat,
    empty_document,
)
from app.parsing.types import ConvertedDocument, ConvertOptions
from tests.unit.parsing.factories import PdfPage, PdfText, make_pdf

OPTS = ConvertOptions()
_CFB = b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1"
_BAD_ZIP = b"PK\x03\x04" + b"reszta nie jest poprawnym zipem " * 10


def _png_bytes() -> bytes:
    buffer = io.BytesIO()
    Image.new("RGB", (2, 2), "blue").save(buffer, format="PNG")
    return buffer.getvalue()


def _jpeg_bytes() -> bytes:
    buffer = io.BytesIO()
    Image.new("RGB", (2, 2), "red").save(buffer, format="JPEG")
    return buffer.getvalue()


def _docx_bytes() -> bytes:
    document = Document()
    document.add_paragraph("Zażółć gęślą jaźń")
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def _xlsx_bytes() -> bytes:
    wb = openpyxl.Workbook()
    wb.active["A1"] = "Kolumna"
    wb.active["A2"] = 3.5
    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


def _pptx_bytes() -> bytes:
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[1])
    slide.shapes.title.text = "Tytuł"
    box = slide.shapes.add_textbox(Emu(0), Emu(0), Emu(4_000_000), Emu(2_000_000))
    box.text_frame.text = "Akapit"
    buffer = io.BytesIO()
    prs.save(buffer)
    return buffer.getvalue()


def _plain_zip_bytes() -> bytes:
    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, "w") as zf:
        zf.writestr("a.txt", "zwykły tekst")
    return buffer.getvalue()


def _vba_zip_bytes() -> bytes:
    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, "w") as zf:
        zf.writestr("xl/workbook.xml", "<workbook/>")
        zf.writestr("xl/vbaProject.bin", "makra")
    return buffer.getvalue()


def _pdf_bytes() -> bytes:
    return make_pdf([PdfPage(texts=[PdfText(72.0, 72.0, 12.0, "Witaj swiecie")])])


def _boom(path: Path, opts: ConvertOptions) -> ConvertedDocument:
    """Konwerter zastępczy rzucający wyjątek ogólny (test opakowania błędów)."""
    raise ValueError("x")


def _empty(path: Path, opts: ConvertOptions) -> ConvertedDocument:
    """Konwerter zastępczy rzucający EmptyDocument (test przepuszczenia bez zmian)."""
    raise empty_document("brak tekstu")


def test_supported_formats_exact() -> None:
    assert registry.supported_formats() == {
        "pdf",
        "docx",
        "doc",
        "xlsx",
        "xlsm",
        "xls",
        "csv",
        "tsv",
        "pptx",
        "ppt",
        "md",
        "html",
        "txt",
        "code",
        "json",
        "jsonl",
        "eml",
        "png",
        "jpg",
        "webp",
        "tiff",
        "bmp",
    }


def test_registry_one_picklable_spec_per_key() -> None:
    for key, specs in registry.REGISTRY.items():
        assert len(specs) == 1, key
        spec = specs[0]
        assert spec.name.startswith(f"{key}."), key
        # round-trip własnej funkcji (noqa w tej samej linii — wymóg ruff)
        assert pickle.loads(pickle.dumps(spec.fn)) is spec.fn, key  # noqa: S301


def test_needs_binaries() -> None:
    assert registry.REGISTRY["doc"][0].needs_binaries == ("wvHtml",)
    assert registry.REGISTRY["ppt"][0].needs_binaries == ("catppt",)
    assert registry.REGISTRY["pdf"][0].needs_binaries == ()
    assert registry.REGISTRY["png"][0].needs_binaries == ()


def test_extension_aliases_derived_from_registry() -> None:
    derived = {ext: key for key, specs in registry.REGISTRY.items() for ext in specs[0].extensions}
    assert derived == registry.EXTENSION_ALIASES
    assert registry.EXTENSION_ALIASES["markdown"] == "md"
    assert registry.EXTENSION_ALIASES["htm"] == "html"
    assert registry.EXTENSION_ALIASES["text"] == "txt"
    assert registry.EXTENSION_ALIASES["ndjson"] == "jsonl"
    assert registry.EXTENSION_ALIASES["jpeg"] == "jpg"
    assert registry.EXTENSION_ALIASES["tif"] == "tiff"


# (treść, zadeklarowane rozszerzenie, oczekiwany klucz albo klasa wyjątku) — magic bytes, krok 2 C9
_DETECT_MAGIC_CASES = [
    pytest.param(_pdf_bytes(), ".txt", "pdf", id="pdf-magic-in-txt"),
    pytest.param(_png_bytes(), ".pdf", "png", id="png-magic-in-pdf"),
    pytest.param(_jpeg_bytes(), ".png", "jpg", id="jpeg-magic"),
    pytest.param(b"RIFF\0\0\0\0WEBP", ".webp", "webp", id="webp-magic"),
    pytest.param(b"II*\x00\x08\x00\x00\x00", ".tif", "tiff", id="tiff-ii"),
    pytest.param(b"MM\x00*\x00\x08\x00\x00", ".tiff", "tiff", id="tiff-mm"),
    pytest.param(b"BM\x8a\x00\x00\x00\x00\x00\x00\x00", ".bmp", "bmp", id="bm-magic-bmp"),
    pytest.param(b"BM tekst zamiast bitmapy", ".txt", "txt", id="bm-in-txt"),
    pytest.param(_docx_bytes(), ".bin", "docx", id="docx-in-bin"),
    pytest.param(_xlsx_bytes(), ".bin", "xlsx", id="xlsx-in-bin"),
    pytest.param(_pptx_bytes(), ".bin", "pptx", id="pptx-in-bin"),
    pytest.param(_xlsx_bytes(), ".xlsm", "xlsm", id="xlsx-as-xlsm"),
    pytest.param(_vba_zip_bytes(), ".bin", "xlsm", id="vba-zip-in-bin"),
    pytest.param(_plain_zip_bytes(), ".zip", UnsupportedFormat, id="plain-zip"),
    pytest.param(_BAD_ZIP, ".docx", "docx", id="broken-zip-docx"),
    pytest.param(_BAD_ZIP, ".zip", UnsupportedFormat, id="broken-zip-zip"),
    pytest.param(_CFB + b"ole-tresc", ".doc", "doc", id="cfb-doc"),
    pytest.param(_CFB + b"ole-tresc", ".xls", "xls", id="cfb-xls"),
    pytest.param(_CFB + b"ole-tresc", ".ppt", "ppt", id="cfb-ppt"),
    pytest.param(_CFB + b"ole-tresc", ".docx", "docx", id="cfb-docx"),
    pytest.param(_CFB + b"ole-tresc", ".bin", UnsupportedFormat, id="cfb-bin"),
    pytest.param(_jpeg_bytes(), ".jpeg", "jpg", id="alias-jpeg"),
    pytest.param(b"II*\x00\x08\x00\x00\x00", ".tif", "tiff", id="alias-tif"),
]

# kroki 1 i 3 C9: brak magic oraz plik 0 B
_DETECT_NOMAGIC_CASES = [
    pytest.param(b"hello world, not a pdf", "pdf", CorruptedFile, id="nomagic-pdf"),
    pytest.param(b"hello", ".jpeg", CorruptedFile, id="nomagic-jpeg"),
    pytest.param(b"hello", ".tif", CorruptedFile, id="nomagic-tif"),
    pytest.param(b"a;b\n1;2", ".CSV", "csv", id="alias-csv-upper"),
    pytest.param(b"# Tytul", ".MD", "md", id="alias-md-upper"),
    pytest.param(b"<p>x</p>", ".htm", "html", id="alias-htm"),
    pytest.param(b"## Naglowek", ".markdown", "md", id="alias-markdown"),
    pytest.param(b"def f():\n    return 1", ".py", "code", id="code-ext"),
    pytest.param(b'{"a": 1}\n', ".ndjson", "jsonl", id="alias-ndjson"),
    pytest.param(b"cokolwiek", ".xyz", UnsupportedFormat, id="unknown-ext"),
    pytest.param(bytes(range(256)) * 4, ".bin", UnsupportedFormat, id="binary-unknown-ext"),
    pytest.param(b"", ".txt", "txt", id="zero-bytes-txt"),
    pytest.param(b"", ".xyz", UnsupportedFormat, id="zero-bytes-xyz"),
]


@pytest.mark.parametrize(("content", "ext", "expected"), _DETECT_MAGIC_CASES)
def test_detect_format_magic(
    content: bytes, ext: str, expected: str | type[Exception], tmp_path: Path
) -> None:
    _assert_detect(tmp_path, content, ext, expected)


@pytest.mark.parametrize(("content", "ext", "expected"), _DETECT_NOMAGIC_CASES)
def test_detect_format_no_magic(
    content: bytes, ext: str, expected: str | type[Exception], tmp_path: Path
) -> None:
    _assert_detect(tmp_path, content, ext, expected)


def _assert_detect(tmp_path: Path, content: bytes, ext: str, expected: object) -> None:
    path = tmp_path / f"plik{ext}"
    path.write_bytes(content)
    if isinstance(expected, str):
        assert registry.detect_format(path, ext) == expected
    else:
        with pytest.raises(expected):  # type: ignore[operator]
            registry.detect_format(path, ext)


def test_detect_pdf_in_txt_exact_brief_example(tmp_path: Path) -> None:
    path = tmp_path / "plik.txt"
    path.write_bytes(make_pdf([PdfPage()]))
    assert registry.detect_format(path, ".txt") == "pdf"


def test_detect_wrong_extension_pdf_message(tmp_path: Path) -> None:
    path = tmp_path / "plik.pdf"
    path.write_bytes(b"hello world, not a pdf")
    with pytest.raises(CorruptedFile) as excinfo:
        registry.detect_format(path, "pdf")
    assert str(excinfo.value) == (
        "File is corrupted or unreadable (.pdf): file content does not match its extension"
    )


def test_detect_unknown_bin_message(tmp_path: Path) -> None:
    path = tmp_path / "plik.bin"
    path.write_bytes(bytes(range(256)) * 4)
    with pytest.raises(UnsupportedFormat) as excinfo:
        registry.detect_format(path, ".bin")
    assert str(excinfo.value) == "Unsupported file format: .bin"


def test_convert_unknown_format(tmp_path: Path) -> None:
    with pytest.raises(UnsupportedFormat) as excinfo:
        registry.convert(tmp_path / "x.txt", "nieznany", OPTS)
    assert str(excinfo.value) == "Unsupported file format: nieznany"


def test_convert_wraps_generic_exception(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    monkeypatch.setitem(
        registry.REGISTRY,
        "txt",
        [registry.ConverterSpec(name="txt.fake", extensions=("txt",), fn=_boom)],
    )
    path = tmp_path / "x.txt"
    path.write_text("treść", encoding="utf-8")
    with pytest.raises(CorruptedFile) as excinfo:
        registry.convert(path, "txt", OPTS)
    assert "ValueError" in str(excinfo.value)
    assert isinstance(excinfo.value.__cause__, ValueError)


def test_convert_passes_empty_document(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    monkeypatch.setitem(
        registry.REGISTRY,
        "txt",
        [registry.ConverterSpec(name="txt.fake", extensions=("txt",), fn=_empty)],
    )
    path = tmp_path / "x.txt"
    path.write_text("treść", encoding="utf-8")
    with pytest.raises(EmptyDocument):
        registry.convert(path, "txt", OPTS)


_SMOKE_CASES = [
    pytest.param(".txt", b"Czesc swiecie\n\nDrugi akapit", id="txt"),
    pytest.param(".md", b"# Tytul\n\nAkapit markdown", id="md"),
    pytest.param(".csv", b"a;b\n1;2", id="csv"),
    pytest.param(".json", b'[{"a": 1, "b": "x"}]', id="json"),
    pytest.param(".html", b"<p>Tresc <strong>html</strong></p>", id="html"),
    pytest.param(".docx", _docx_bytes(), id="docx"),
    pytest.param(".xlsx", _xlsx_bytes(), id="xlsx"),
    pytest.param(".pptx", _pptx_bytes(), id="pptx"),
    pytest.param(".pdf", _pdf_bytes(), id="pdf"),
    pytest.param(".png", _png_bytes(), id="png"),
]


@pytest.mark.parametrize(("ext", "content"), _SMOKE_CASES)
def test_smoke_detect_then_convert(ext: str, content: bytes, tmp_path: Path) -> None:
    path = tmp_path / f"plik{ext}"
    path.write_bytes(content)
    fmt = registry.detect_format(path, path.suffix)
    opts = ConvertOptions(ocr_available=True) if ext == ".png" else OPTS
    doc = registry.convert(path, fmt, opts)
    assert isinstance(doc, ConvertedDocument)
    assert doc.html or doc.pending_pages
    if ext == ".png":
        assert doc.pending_pages
