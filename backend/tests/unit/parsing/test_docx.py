"""Testy konwerterów DOCX i DOC (T5, kontrakt C8)."""

import io
import pickle
import shutil
import subprocess
from pathlib import Path
from types import SimpleNamespace

import pytest
from docx import Document
from PIL import Image

from app.parsing.converters.doc import convert_doc
from app.parsing.converters.docx import convert_docx
from app.parsing.errors import (
    BinaryMissing,
    CorruptedFile,
    EmptyDocument,
    EncryptedFile,
)
from app.parsing.types import ConvertOptions

_CFB = b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1"


def _docx(path: Path, *, image_only: bool = False) -> Path:
    document = Document()
    if image_only:
        buffer = io.BytesIO()
        Image.new("RGB", (20, 20), "red").save(buffer, format="PNG")
        document.add_picture(io.BytesIO(buffer.getvalue()))
    else:
        document.add_paragraph("Raport kwartalny", style="Title")
        document.add_paragraph("Cel", style="Heading 1")
        document.add_paragraph("Szczegóły", style="Heading 2")
        document.add_paragraph("Zażółć gęślą jaźń")
        document.add_paragraph("Drugi akapit")
        document.add_paragraph("Punkt", style="List Bullet")
        table = document.add_table(rows=3, cols=2)
        for row_nr, row in enumerate(table.rows):
            for col_nr, cell in enumerate(row.cells):
                cell.text = f"Nagłówek {col_nr}" if row_nr == 0 else f"wiersz {row_nr}{col_nr}"
    document.core_properties.title = "Raport"
    document.core_properties.author = "Jan"
    document.save(path)
    return path


def test_docx_full_document(tmp_path: Path) -> None:
    result = convert_docx(_docx(tmp_path / "doc.docx"), ConvertOptions())
    assert "<h1>" in result.html
    assert "<h2>" in result.html
    assert "<p>" in result.html
    assert "<table><thead><tr><th>" in result.html
    assert "Zażółć gęślą jaźń" in result.html
    assert result.metadata["title"] == "Raport"
    assert result.metadata["author"] == "Jan"


def test_docx_heading_then_paragraph(tmp_path: Path) -> None:
    document = Document()
    document.add_paragraph("Cel", style="Heading 1")
    document.add_paragraph("Tekst")
    document.save(tmp_path / "mini.docx")
    result = convert_docx(tmp_path / "mini.docx", ConvertOptions())
    assert "<h1>Cel</h1><p>Tekst</p>" in result.html


def test_docx_image_only_raises_empty(tmp_path: Path) -> None:
    with pytest.raises(EmptyDocument) as excinfo:
        convert_docx(_docx(tmp_path / "img.docx", image_only=True), ConvertOptions())
    assert "document has no text" in str(excinfo.value)


def test_docx_encrypted(tmp_path: Path) -> None:
    path = tmp_path / "enc.docx"
    path.write_bytes(_CFB + b"\0" * 504)
    with pytest.raises(EncryptedFile):
        convert_docx(path, ConvertOptions())


def test_docx_corrupted(tmp_path: Path) -> None:
    path = tmp_path / "junk.docx"
    path.write_bytes(b"to nie jest archiwum ZIP")
    with pytest.raises(CorruptedFile) as excinfo:
        convert_docx(path, ConvertOptions())
    assert excinfo.value.details["format"] == "docx"
    assert excinfo.value.details["reason"] == "not a valid ZIP container"


_WV_OUT = (
    "<html><body><h1>T</h1><p>Treść</p>"
    "<table><tr><td>a</td></tr><tr><td>1</td></tr></table></body></html>"
)


def test_doc_binary_missing(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    monkeypatch.setattr(shutil, "which", lambda _name: None)
    with pytest.raises(BinaryMissing) as excinfo:
        convert_doc(tmp_path / "stary.doc", ConvertOptions())
    assert excinfo.value.details["binary"] == "wvHtml"


def test_doc_success(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    source = tmp_path / "stary.doc"
    source.write_bytes(b"binaryzny word 97")

    def fake_run(args: list[str], **_kw: object) -> SimpleNamespace:
        targetdir = Path(next(a for a in args if a.startswith("--targetdir=")).split("=", 1)[1])
        (targetdir / "out.html").write_text(_WV_OUT, encoding="utf-8")
        return SimpleNamespace(returncode=0, stdout=b"", stderr=b"")

    monkeypatch.setattr(shutil, "which", lambda _name: "/usr/bin/wvHtml")
    monkeypatch.setattr(subprocess, "run", fake_run)
    result = convert_doc(source, ConvertOptions())
    assert "<h1>T</h1>" in result.html
    assert "<p>Treść</p>" in result.html
    assert "<thead>" in result.html


def test_doc_wvhtml_failure(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    monkeypatch.setattr(shutil, "which", lambda _name: "/usr/bin/wvHtml")
    proc = SimpleNamespace(returncode=1, stdout=b"", stderr=b"boom")
    monkeypatch.setattr(subprocess, "run", lambda *_a, **_kw: proc)
    with pytest.raises(CorruptedFile) as excinfo:
        convert_doc(tmp_path / "stary.doc", ConvertOptions())
    assert excinfo.value.details["format"] == "doc"


def test_converters_picklable() -> None:
    assert pickle.dumps(convert_docx)
    assert pickle.dumps(convert_doc)
