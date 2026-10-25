"""Testy T3 (E5-H6d): obrazy slajdów PPTX → fake-vision przez `complete_pages` (E2E wiązania)."""

import hashlib
import io
import os
from pathlib import Path

import pytest
from PIL import Image
from pptx import Presentation
from pptx.util import Emu
from pydantic import SecretStr

from app.core.config import Settings
from app.llm.openai_gateway import OpenAIGateway
from app.llm.types import ModelRef
from app.parsing.converters.pptx import convert_pptx
from app.parsing.recognizers import VlmRecognizer, complete_pages
from app.parsing.types import ConvertOptions

pytestmark = [
    pytest.mark.fakellm,
    pytest.mark.skipif(
        not os.environ.get("WD_LLM_BASE_URL"),
        reason="fakellm: brak WD_LLM_BASE_URL (eksportuje verify.sh)",
    ),
]


def _png(size: int) -> io.BytesIO:
    buffer = io.BytesIO()
    Image.new("RGB", (size, size), "blue").save(buffer, format="PNG")
    buffer.seek(0)
    return buffer


def _image_only_deck(path: Path) -> Path:
    """Slajd bez tekstu, wyłącznie z obrazem (sekcja `<section data-slide="1">` pusta)."""
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    slide.shapes.add_picture(_png(800), Emu(0), Emu(0), Emu(8000000), Emu(6000000))
    prs.save(path)
    return path


def _text_image_deck(path: Path) -> Path:
    """Slajd z tekstem i dużym obrazem (nad progiem) — ścieżka `reason="slide image"`."""
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    slide.shapes.add_textbox(
        Emu(0), Emu(0), Emu(6000000), Emu(1000000)
    ).text_frame.text = "Slajd z obrazem i tekstem."
    slide.shapes.add_picture(_png(800), Emu(0), Emu(0), Emu(8000000), Emu(6000000))
    prs.save(path)
    return path


def _gateway() -> OpenAIGateway:
    return OpenAIGateway(
        Settings(
            env="test",
            llm_base_url=os.environ["WD_LLM_BASE_URL"],
            llm_api_key=SecretStr("fake"),
            llm_max_retries=0,
        )
    )


async def test_pptx_image_only_slide_fakellm(tmp_path: Path) -> None:
    """Slajd image-only: fragment fake-vision ląduje w `<section data-slide="1">`."""
    doc = convert_pptx(
        _image_only_deck(tmp_path / "obraz.pptx"), ConvertOptions(vision_available=True)
    )
    (pending,) = doc.pending_pages
    gateway = _gateway()
    try:
        vlm = VlmRecognizer(gateway, ModelRef(None, "fake-vision", "image2text"))
        result = await complete_pages(doc, ocr=None, vision=vlm)
    finally:
        await gateway.aclose()
    assert result.html == (
        '<section data-slide="1"><p>OCR:'
        + hashlib.sha256(pending.image_png).hexdigest()[:8]
        + "</p></section>"
    )
    assert result.stats["pages_vision"] == 1
    assert result.pending_pages == []


async def test_pptx_text_slide_image_fakellm(tmp_path: Path) -> None:
    """Slajd z tekstem i dużym obrazem: fragment fake-vision ląduje w sekcji slajdu (CONTRACT p. 2).

    T5: `_insert_fragment` wstawia fragment przed `</section>` istniejącej sekcji `data-slide="1"`;
    dokument ma dokładnie jedną sekcję slajdu i żadnej `data-page`.
    """
    doc = convert_pptx(
        _text_image_deck(tmp_path / "tekst-obraz.pptx"), ConvertOptions(vision_available=True)
    )
    (pending,) = doc.pending_pages
    assert pending.reason == "slide image"
    gateway = _gateway()
    try:
        vlm = VlmRecognizer(gateway, ModelRef(None, "fake-vision", "image2text"))
        result = await complete_pages(doc, ocr=None, vision=vlm)
    finally:
        await gateway.aclose()
    fragment = "<p>OCR:" + hashlib.sha256(pending.image_png).hexdigest()[:8] + "</p>"
    assert result.html.count('<section data-slide="1">') == 1
    assert "<section data-page=" not in result.html
    assert fragment in result.html
    assert result.html.index(fragment) < result.html.index("</section>")
    assert result.stats["pages_vision"] == 1
    assert result.pending_pages == []
