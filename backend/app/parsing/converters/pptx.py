"""Konwertery PPTX i PPT (kontrakt C8) — slajdy, tabele, notatki, obrazy do OCR/VLM."""

from __future__ import annotations

import io
from collections.abc import Iterator
from pathlib import Path
from typing import Literal, cast

from PIL import Image
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE
from pptx.shapes.autoshape import Shape
from pptx.shapes.base import BaseShape
from pptx.shapes.graphfrm import GraphicFrame
from pptx.shapes.group import GroupShape
from pptx.shapes.picture import Picture
from pptx.shapes.shapetree import GroupShapes, SlideShapes
from pptx.slide import Slide

from app.parsing.converters.common import check_zip, finish, png_bytes, run_binary
from app.parsing.encoding import decode_bytes
from app.parsing.errors import empty_document
from app.parsing.html import esc, paragraphs_html, section_html, table_html, text_chars
from app.parsing.types import ConvertedDocument, ConvertOptions, PendingPage

# Próg „obraz slajdu": ułamek pola slajdu (EMU^2 obrazu / EMU^2 slajdu), DEC-264.
# Poniżej progu (ikony, logo, ozdoby) obraz nie jest odsyłany do OCR/VLM.
_MIN_IMAGE_EMU2 = 0.05


def convert_pptx(path: Path, opts: ConvertOptions) -> ConvertedDocument:
    """PPTX → sekcje `data-slide`: tytuł `<h2>`, akapity, tabele, notatki; obrazy do OCR/VLM."""
    check_zip(path, "pptx", opts)
    prs = Presentation(str(path))
    slide_area = (prs.slide_width or 0) * (prs.slide_height or 0)
    sections: list[str] = []
    pending: list[PendingPage] = []
    warnings: list[str] = []
    for number, slide in enumerate(prs.slides, start=1):
        inner = _slide_inner(slide, number, opts, pending, warnings, slide_area=slide_area)
        sections.append(section_html("slide", number, inner))
    return finish(
        "".join(sections),
        diagnosis="presentation has no text",
        pages=len(prs.slides),
        pending=pending,
        warnings=warnings,
    )


def _slide_inner(  # noqa: PLR0913 — 6. argument (pole slajdu) wymaga próg obrazów DEC-264
    slide: Slide,
    number: int,
    opts: ConvertOptions,
    pending: list[PendingPage],
    warnings: list[str],
    *,
    slide_area: int,
) -> str:
    """Wnętrze sekcji slajdu: tytuł, kształty (rekurencyjnie), notatki.

    Slajd z tekstem (a bez tytułu) dostaje `<h2>Slajd N</h2>`, by chunker `naive` (granica na
    nagłówku) nie łączył slajdów. Slajd bez tekstu (pusty, sam obraz) nie dostaje etykiety —
    inaczej `text_chars` uznałby ją za treść i `EmptyDocument` przestałby działać (DEC-263).
    """
    title_text, body, images = _slide_parts(slide)
    if title_text is None and not body and images:
        _handle_image_only(number, opts, pending, warnings, max(images, key=_area)[1])
    elif images:
        # Obraz na slajdzie z tekstem: do rozpoznania tylko powyżej progu (ikony/ozdoby odpada).
        big = [candidate for candidate in images if candidate[0] >= _MIN_IMAGE_EMU2 * slide_area]
        if big:
            _handle_slide_image(number, opts, pending, warnings, max(big, key=_area)[1])
    if slide.has_notes_slide:
        notes_frame = slide.notes_slide.notes_text_frame
        if notes_frame is not None and notes_frame.text.strip():
            body.append(f"<aside data-notes>{paragraphs_html(notes_frame.text)}</aside>")
    if title_text is not None:
        body.insert(0, f"<h2>{esc(title_text)}</h2>")
    elif body:
        body.insert(0, f"<h2>{esc(f'Slajd {number}')}</h2>")
    return "".join(body)


def _slide_parts(slide: Slide) -> tuple[str | None, list[str], list[tuple[int, Picture]]]:
    """Części slajdu: tytuł, treść (`<p>`/tabela) i obrazy wraz z polem kształtu (EMU^2)."""
    title = slide.shapes.title
    title_text: str | None = None
    if title is not None and title.has_text_frame and title.text.strip():
        title_text = title.text
    body: list[str] = []
    images: list[tuple[int, Picture]] = []
    for shape in _walk(slide.shapes):
        if title is not None and shape.shape_id == title.shape_id:
            continue  # tytuł już jest w <h2> — nie powtarzamy go jako <p>
        if shape.has_text_frame:
            for paragraph in cast(Shape, shape).text_frame.paragraphs:
                if paragraph.text.strip():
                    body.append(f"<p>{esc(paragraph.text)}</p>")
        elif shape.has_table:
            rows = [
                [cell.text for cell in row.cells] for row in cast(GraphicFrame, shape).table.rows
            ]
            body.append(table_html(rows))
        elif shape.shape_type == MSO_SHAPE_TYPE.PICTURE:
            images.append((shape.width * shape.height, cast(Picture, shape)))
    return title_text, body, images


def _area(candidate: tuple[int, Picture]) -> int:
    """Klucz wyboru „największego” obrazu: pole kształtu (width * height)."""
    return candidate[0]


def _walk(shapes: SlideShapes | GroupShapes) -> Iterator[BaseShape]:
    """Kształty slajdu z rozwinięciem grup (rekurencja po `shape.shapes`)."""
    for shape in shapes:
        if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
            yield from _walk(cast(GroupShape, shape).shapes)
        else:
            yield shape


def _handle_image_only(
    number: int,
    opts: ConvertOptions,
    pending: list[PendingPage],
    warnings: list[str],
    picture: Picture,
) -> None:
    """Slajd bez tekstu z obrazem: PendingPage do OCR/VLM albo warning."""
    if not (opts.vision_available or opts.ocr_available):
        warnings.append(f"slide {number}: image only, no OCR/VLM available")
        return
    image = _decode(picture)
    if image is None:
        warnings.append(f"slide {number}: image could not be decoded")
        return
    mode: Literal["ocr", "vision"] = "vision" if opts.vision_available else "ocr"
    pending.append(
        PendingPage(page=number, mode=mode, image_png=png_bytes(image), reason="image-only slide")
    )


def _handle_slide_image(
    number: int,
    opts: ConvertOptions,
    pending: list[PendingPage],
    warnings: list[str],
    picture: Picture,
) -> None:
    """Obraz powyżej progu na slajdzie z tekstem: PendingPage do OCR/VLM albo warning (DEC-264)."""
    if not (opts.vision_available or opts.ocr_available):
        warnings.append(f"slide {number}: image(s) not recognized, no OCR/VLM available")
        return
    image = _decode(picture)
    if image is None:
        warnings.append(f"slide {number}: image could not be decoded")
        return
    mode: Literal["ocr", "vision"] = "vision" if opts.vision_available else "ocr"
    pending.append(
        PendingPage(page=number, mode=mode, image_png=png_bytes(image), reason="slide image")
    )


def _decode(picture: Picture) -> Image.Image | None:
    """Jedna próba dekodowania obrazu kształtu; `None`, gdy PIL go nie rozpoznaje."""
    try:
        image = Image.open(io.BytesIO(picture.image.blob))
        image.load()
    except (OSError, ValueError):
        return None
    return image


def convert_ppt(path: Path, opts: ConvertOptions) -> ConvertedDocument:
    """Starszy .ppt → tekst przez catppt; części po `\f` stają się sekcjami `data-slide`."""
    stdout = run_binary("catppt", [str(path)], fmt="ppt")
    text, _encoding = decode_bytes(stdout, fmt="ppt")
    sections = [
        section_html("slide", number, paragraphs_html(part))
        for number, part in enumerate((p for p in text.split("\f") if p.strip()), start=1)
    ]
    html = "".join(sections)
    if text_chars(html) < opts.min_chars:
        raise empty_document("legacy .ppt contains only images/OLE objects")
    return finish(html, diagnosis="presentation has no text", pages=len(sections))
