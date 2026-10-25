"""Pomocniki testów pipeline'u (brief T7): budowa plików, kontener, kontekst zadania."""

import io
import logging
from collections.abc import AsyncIterator
from pathlib import Path

from docx import Document as DocxDocument
from openpyxl import Workbook
from sqlalchemy.ext.asyncio import AsyncSession, async_sessionmaker

from app.core.config import Settings
from app.core.container import Container
from app.db.repositories import models as models_repo
from app.db.uow import transaction
from app.services.ingest import UploadSource
from app.worker.context import TaskContext
from app.worker.queue import DbTaskQueue, TaskRow
from tests.integration.ingest.conftest import TIME_MS
from tests.unit.parsing.factories import PdfPage, PdfText, make_pdf

WORKER_ID = "w:1:0"

# Treść strony PDF (>= 150 znaków — strona tekstowa przy strategii auto, 05 §8.2).
PDF_PAGE_TEXT = "Lorem ipsum dolor sit amet, consectetur adipiscing elit. " * 6


async def _aiter(data: bytes) -> AsyncIterator[bytes]:
    """Strumień bajtów jako async generator (UploadSource)."""
    yield data


def make_source(content: bytes, filename: str) -> UploadSource:
    """UploadSource z bajtów i nazwą pliku."""
    return UploadSource(stream=_aiter(content), filename=filename)


def make_pdf_bytes(pages: int = 2) -> bytes:
    """PDF z `pages` stronami tekstu (fabryka E1-4, tests/unit/parsing/factories)."""
    return make_pdf(
        [
            PdfPage(texts=[PdfText(x=50, top=50, size=12, text=PDF_PAGE_TEXT)]),
        ]
        * pages
    )


def make_docx_bytes() -> bytes:
    """DOCX: nagłówek + 3 akapity (python-docx)."""
    document = DocxDocument()
    document.add_heading("Umowa najmu", level=1)
    for i in range(3):
        document.add_paragraph(
            f"Akapit numer {i}: lokator zobowiązuje się uiszczać czynsz do piątego dnia miesiąca."
        )
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def make_xlsx_bytes() -> bytes:
    """XLSX: wiersz nagłówkowy + 5 wierszy danych (openpyxl)."""
    workbook = Workbook()
    sheet = workbook.active
    assert sheet is not None  # aktywny arkusz zawsze istnieje
    sheet.append(["Produkt", "Cena"])
    for i in range(5):
        sheet.append([f"produkt-{i}", i + 1])
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


async def register_models(sm: async_sessionmaker[AsyncSession]) -> None:
    """Rejestruje modele fake-llm w bazie: fake-embed-1024 (embedding) i fake-chat (chat)."""
    async with transaction(sm) as session:
        await models_repo.insert(
            session,
            id="e" * 32,
            name="fake-embed-1024",
            model_type="embedding",
            is_default=True,
            now_ms=TIME_MS,
        )
        await models_repo.insert(
            session,
            id="c" * 32,
            name="fake-chat",
            model_type="chat",
            is_default=True,
            now_ms=TIME_MS,
        )


async def make_container(
    migrated_db: Settings,
    tmp_path: Path,
    *,
    ocr_engine: str = "none",
    settings_overrides: dict[str, object] | None = None,
) -> Container:
    """Kontener testowy: storage w tmp_path, ocr_engine podany; start + rejestr modeli.

    ``settings_overrides`` (E5-H8b T3) pozwala nadpisać pojedyncze pola ``Settings``
    (np. ``embed_batch_size``) bez zmiany zachowania dotychczasowych wywołań.
    """
    update: dict[str, object] = {
        "storage_root": str(tmp_path),
        "ocr_engine": ocr_engine,
        **(settings_overrides or {}),
    }
    container = Container(migrated_db.model_copy(update=update))
    await container.start()
    await register_models(container.sessionmaker)
    return container


def worker_queue(container: Container) -> DbTaskQueue:
    """Kolejka zadaniowa na usługach kontenera (claim/complete/request_cancel)."""
    assert container.sessionmaker is not None
    return DbTaskQueue(container.sessionmaker, container.clock, container.id_gen)


def build_context(container: Container, task: TaskRow, queue: DbTaskQueue) -> TaskContext:
    """TaskContext jak u workera: task, worker_id, kontener, logger, kolejka."""
    logger: logging.LoggerAdapter[logging.Logger] = logging.LoggerAdapter(
        logging.getLogger("test-pipeline"), {}
    )
    return TaskContext(task, WORKER_ID, container, logger, queue)


async def claim_context(
    container: Container, document_id: str, *, clear_existing: bool = False
) -> TaskContext:
    """Prawdziwa kolejka: enqueue_parse → claim → TaskContext (brief T7)."""
    await container.ingest.enqueue_parse([document_id], clear_existing=clear_existing)
    task = await worker_queue(container).claim(WORKER_ID, ["parse_document"])
    assert task is not None, "zadanie musi być do wzięcia"
    return build_context(container, task, worker_queue(container))
