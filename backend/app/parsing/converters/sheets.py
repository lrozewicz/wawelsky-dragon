"""Konwertery arkuszy: xlsx/xlsm (openpyxl), xls (xlrd), csv/tsv (kontrakt C8)."""

from __future__ import annotations

import csv
import datetime as dt
import io
from collections.abc import Iterable, Iterator, Sequence
from pathlib import Path
from typing import Any

import openpyxl
import xlrd
from xlrd.xldate import xldate_as_datetime

from app.parsing.converters.common import check_zip, finish, read_bytes
from app.parsing.encoding import decode_bytes
from app.parsing.errors import corrupted_file, encrypted_file, limit_exceeded
from app.parsing.html import section_html, table_html
from app.parsing.types import ConvertedDocument, ConvertOptions

_DIAGNOSIS = "spreadsheet has no data"
_HIDDEN_WARNING = "sheet '{name}' is hidden, skipped"


class _SemicolonDialect(csv.excel):
    """Fallback dialektu CSV: średnik (pierwsza linia ma więcej ';' niż ',')."""

    delimiter = ";"


class _SheetRows:
    """Wiersze jednego arkusza: pomijanie pustych wierszy, obcinanie pustych
    kolumn na końcu (po całym arkuszu) i wspólny limit wierszy dokumentu (C8)."""

    def __init__(self, opts: ConvertOptions, used: int = 0) -> None:
        self._opts = opts
        self.used = used  # wiersze policzone dotąd w całym dokumencie
        self.rows: list[list[str]] = []

    def feed(self, raw_rows: Iterable[Sequence[object]]) -> None:
        """Dodaje wiersze; limit sprawdzany natychmiast po przekroczeniu."""
        for raw in raw_rows:
            cells = [cell_str(value) for value in raw]
            if not any(cells):
                continue  # pusty wiersz
            self.used += 1
            if self.used > self._opts.max_sheet_rows:
                raise limit_exceeded("sheet rows", self.used, self._opts.max_sheet_rows)
            self.rows.append(cells)
        self._trim_trailing_columns()

    def _trim_trailing_columns(self) -> None:
        """Obcina końcowe kolumny puste w każdym zachowanym wierszu."""
        width = 0
        for row in self.rows:
            for index in range(len(row) - 1, -1, -1):
                if row[index]:
                    width = max(width, index + 1)
                    break
        self.rows = [row[:width] for row in self.rows]


def cell_str(value: object) -> str:
    """Komórka arkusza → tekst (C8): bool TRUE/FALSE, daty ISO, całkowity float bez .0."""
    if value is None:
        return ""
    if isinstance(value, bool):
        return "TRUE" if value else "FALSE"
    if isinstance(value, dt.datetime):
        return value.date().isoformat() if value.time() == dt.time() else value.isoformat()
    if isinstance(value, (dt.date, dt.time)):
        return value.isoformat()
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value)


def _finish_sheets(sections: list[str], warnings: list[str]) -> ConvertedDocument:
    """Skleja sekcje arkuszy i zamyka dokument wspólną diagnozą (C8)."""
    return finish("".join(sections), diagnosis=_DIAGNOSIS, warnings=warnings)


def convert_xlsx(path: Path, opts: ConvertOptions) -> ConvertedDocument:
    """XLSX i XLSM (ta sama funkcja) → sekcje <section data-sheet> z tabelami (C8)."""
    fmt = "xlsm" if path.suffix.lower() == ".xlsm" else "xlsx"
    check_zip(path, fmt, opts)
    sections: list[str] = []
    warnings: list[str] = []
    used = 0
    workbook = openpyxl.load_workbook(path, read_only=True, data_only=True)
    try:
        for name in workbook.sheetnames:
            sheet = workbook[name]
            if sheet.sheet_state != "visible":
                warnings.append(_HIDDEN_WARNING.format(name=name))
                continue
            rows = _SheetRows(opts, used)
            rows.feed(sheet.iter_rows(values_only=True))
            used = rows.used
            if rows.rows:  # pusty arkusz pominięty
                sections.append(section_html("sheet", name, table_html(rows.rows)))
    finally:
        workbook.close()
    return _finish_sheets(sections, warnings)


def convert_xls(path: Path, opts: ConvertOptions) -> ConvertedDocument:
    """XLS (xlrd) → sekcje <section data-sheet>; daty przez xldate_as_datetime (C8)."""
    try:
        book = xlrd.open_workbook(str(path), on_demand=True)
    except xlrd.XLRDError as exc:
        if "encrypted" in str(exc).lower():
            raise encrypted_file("xls") from None
        raise corrupted_file("xls", str(exc)) from None
    sections: list[str] = []
    warnings: list[str] = []
    used = 0
    try:
        for index in range(book.nsheets):
            sheet = book.sheet_by_index(index)
            if sheet.visibility != 0:
                warnings.append(_HIDDEN_WARNING.format(name=sheet.name))
                continue
            rows = _SheetRows(opts, used)
            rows.feed(_xls_cells(sheet, book.datemode))
            used = rows.used
            if rows.rows:  # pusty arkusz pominięty
                sections.append(section_html("sheet", sheet.name, table_html(rows.rows)))
    finally:
        book.release_resources()
    return _finish_sheets(sections, warnings)


def _xls_cells(sheet: xlrd.sheet.Sheet, datemode: int) -> Iterator[list[object]]:
    """Wiersze arkusza xlrd; komórki datowe zamieniane na datetime (C8)."""
    for row_index in range(sheet.nrows):
        yield [
            (
                xldate_as_datetime(cell.value, datemode)
                if cell.ctype == xlrd.XL_CELL_DATE
                else cell.value
            )
            for cell in sheet.row(row_index)
        ]


def convert_csv(path: Path, opts: ConvertOptions) -> ConvertedDocument:
    """CSV → jedna tabela HTML; separator sniffowany, fallback ';' lub ',' (C8)."""
    text, encoding = decode_bytes(read_bytes(path), fmt="csv")
    return _delimited_document(text, _sniff_dialect(text), encoding, opts)


def convert_tsv(path: Path, opts: ConvertOptions) -> ConvertedDocument:
    """TSV → jedna tabela HTML; separator tabulator (C8)."""
    text, encoding = decode_bytes(read_bytes(path), fmt="tsv")
    return _delimited_document(text, csv.excel_tab, encoding, opts)


def _sniff_dialect(text: str) -> Any:
    """Dialekt CSV ze Sniffera; po csv.Error fallback na ';' albo ',' (C8)."""
    try:
        return csv.Sniffer().sniff(text[:65536], delimiters=",;\t|")
    except csv.Error:
        first_line = text.split("\n", 1)[0]
        if first_line.count(";") > first_line.count(","):
            return _SemicolonDialect
        return csv.excel


def _delimited_document(
    text: str, dialect: Any, encoding: str, opts: ConvertOptions
) -> ConvertedDocument:
    """Wspólne domknięcie csv/tsv: wiersze przez csv.reader, encoding w metadanych."""
    rows = _SheetRows(opts)
    rows.feed(csv.reader(io.StringIO(text, newline=""), dialect))
    return finish(
        table_html(rows.rows),
        diagnosis=_DIAGNOSIS,
        metadata={"encoding": encoding},
    )
