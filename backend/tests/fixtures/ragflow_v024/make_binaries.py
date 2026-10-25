"""Generator deterministycznych plików binarnych fixture'u RAGFlow v0.24 (brief E4-7a T3).

Tworzy ``raport.pdf`` (reportlab: jedna strona A4 z prawdziwą warstwą tekstową) i
``umowa.docx`` (python-docx: nagłówek + akapity, archiwum ZIP przepisane na stałe znaczniki
czasu). Treść jest wyłącznie syntetyczna — bez danych osobowych. Wynik jest powtarzalny:
dwa uruchomienia dają identyczny SHA-256 obu plików (test
``tests/unit/tools/test_ragflow_fixture_binaries_e4_7a.py``).

Uruchomienie (z katalogu ``backend/``)::

    uv run --frozen python tests/fixtures/ragflow_v024/make_binaries.py

Skrypt wypisuje ``<nazwa> <size> <sha256>`` — wartości do ``data.sql`` i ``expectations.json``.
"""

from __future__ import annotations

import hashlib
import io
import sys
import zipfile
from datetime import datetime
from pathlib import Path

import docx
from reportlab.lib.pagesizes import A4
from reportlab.pdfgen.canvas import Canvas

#: Katalog-blob dokumentów fi01/f02 w fixture'ze (``files/<bucket>/<location>``).
BUCKET = "d1e2f3a4b5c60718293a4b5c6d7e8f01"
FILES_DIR = Path(__file__).resolve().parent / "files" / BUCKET

#: Stały znacznik czasu wpisów ZIP (DOCX) — ``2025-10-01 08:05:00`` jak reszta fixture'u.
ZIP_DATE_TIME = (2025, 10, 1, 8, 5, 0)
#: Stała data ``core_properties`` DOCX (created/modified).
DOCX_TIMESTAMP = datetime(2025, 10, 1, 8, 5, 0)  # data bez strefy, jak reszta fixture'u

#: Treść PDF — bez diakrytyków (standardowa Helvetica reportlabu używa WinAnsiEncoding,
#: więc polskie znaki diakrytyczne nie dałyby się wyekstrahować). Pierwsze zdanie jest
#: sprawdzane słownie w teście.
PDF_LINES = (
    "Raport kwartalny Wawelskiego Smoka",
    "Podsumowanie wynikow indeksowania dokumentow w kwartale.",
    "System przetworzyl wszystkie zbiory bez bledow krytycznych.",
    "Liczba fragmentow wzrosla o dwanascie procent.",
    "Raport przygotowano na potrzeby testow syntetycznych.",
)

#: Nagłówek i akapity DOCX (treść syntetyczna, po polsku).
DOCX_HEADING = "Umowa o świadczenie usług indeksowania — dane syntetyczne"
DOCX_PARAGRAPHS = (
    "Umowa opisuje syntetyczny zbiór dokumentów używany wyłącznie w testach migracji.",
    "Strony ustalają, że indeksowanie obejmuje pliki przekazane w ramach fixture'u.",
    "Umowa nie zawiera danych osobowych ani informacji poufnych.",
)

PDF_NAME = "raport.pdf"
DOCX_NAME = "umowa.docx"


def build_pdf() -> bytes:
    """Bajty jednostronicowego PDF A4 z warstwą tekstową (powtarzalne)."""
    buf = io.BytesIO()
    canvas = Canvas(buf, invariant=1, pagesize=A4)
    y = A4[1] - 96.0
    for line in PDF_LINES:
        canvas.drawString(72.0, y, line)
        y -= 24.0
    canvas.showPage()
    canvas.save()
    return buf.getvalue()


def build_docx() -> bytes:
    """Bajty DOCX z nagłówkiem i akapitami (powtarzalne).

    ``python-docx`` zapisuje archiwum ZIP ze znacznikiem czasu „teraz", więc po zapisie
    przepisujemy je ``zipfile`` ze stałym ``date_time`` każdego wpisu i stałymi atrybutami.
    """
    document = docx.Document()
    document.add_heading(DOCX_HEADING, level=1)
    for paragraph in DOCX_PARAGRAPHS:
        document.add_paragraph(paragraph)
    properties = document.core_properties
    properties.author = "fixture"
    properties.created = DOCX_TIMESTAMP
    properties.modified = DOCX_TIMESTAMP

    buf = io.BytesIO()
    document.save(buf)

    out = io.BytesIO()
    with zipfile.ZipFile(buf) as source, zipfile.ZipFile(out, "w") as target:
        # Kolejność wpisów z pakietu python-docx jest stała (spacer po relacjach pakietu).
        for name in source.namelist():
            info = zipfile.ZipInfo(name, date_time=ZIP_DATE_TIME)
            info.compress_type = zipfile.ZIP_DEFLATED
            info.create_system = 3  # Unix — ten sam bajt niezależnie od platformy
            info.external_attr = 0o600 << 16
            target.writestr(info, source.read(name))
    return out.getvalue()


def main() -> int:
    """Zapisuje oba pliki w ``files/<bucket>/`` i drukuje rozmiar oraz SHA-256."""
    built = {PDF_NAME: build_pdf(), DOCX_NAME: build_docx()}
    FILES_DIR.mkdir(parents=True, exist_ok=True)
    for name, data in built.items():
        (FILES_DIR / name).write_bytes(data)
        print(f"{name} {len(data)} {hashlib.sha256(data).hexdigest()}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
