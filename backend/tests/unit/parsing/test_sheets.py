"""Testy konwerterów arkuszy: xlsx/xlsm, xls, csv/tsv (T6, kontrakt C8)."""

import pickle
import re
from datetime import date, datetime
from pathlib import Path

import openpyxl
import pytest
import xlrd

from app.parsing.converters.sheets import (
    cell_str,
    convert_csv,
    convert_tsv,
    convert_xls,
    convert_xlsx,
)
from app.parsing.errors import (
    CorruptedFile,
    EmptyDocument,
    EncryptedFile,
    LimitExceeded,
)
from app.parsing.types import ConvertOptions

OPTS = ConvertOptions()
FIXTURE_XLS = Path(__file__).parents[2] / "fixtures" / "parsing" / "sample.xls"

_CFB = b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1"  # nagłówek CFB/OLE — zaszyfrowany OOXML


def _compact(html: str) -> str:
    """Porównanie ignorujące białe znaki między tagami (treść tekstowa bez zmian)."""
    return re.sub(r">\s+<", "><", html).strip()


def _make_xlsx(path: Path, *, hidden: bool = True) -> None:
    """Sprzedaż: nagłówek + 2 wiersze z pustym wierszem w środku i pustą kolumną na końcu."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Sprzedaż"
    ws.append(["Produkt", "Cena", None])
    ws.append([None, None, None])
    ws.append(["Ala", "3,14", None])
    ws.append(["Kot", "2,50", None])
    if hidden:
        secret = wb.create_sheet("Ukryty")
        secret.sheet_state = "hidden"
        secret.append(["sekret"])
    wb.save(path)


# ---------------------------------------------------------------- cell_str


@pytest.mark.parametrize(
    ("value", "expected"),
    [
        (None, ""),
        (True, "TRUE"),
        (False, "FALSE"),
        (3.0, "3"),
        (3.5, "3.5"),
        (7, "7"),
        (datetime(2024, 1, 2), "2024-01-02"),
        (datetime(2024, 1, 2, 10, 30), "2024-01-02T10:30:00"),
        (date(2024, 1, 2), "2024-01-02"),
        ("x", "x"),
    ],
)
def test_cell_str(value: object, expected: str) -> None:
    assert cell_str(value) == expected


# ---------------------------------------------------------------- xlsx/xlsm


def test_xlsx_sections_hidden_skipped(tmp_path: Path) -> None:
    path = tmp_path / "book.xlsx"
    _make_xlsx(path)
    doc = convert_xlsx(path, OPTS)
    assert _compact(doc.html) == (
        '<section data-sheet="Sprzedaż"><table><thead><tr>'
        "<th>Produkt</th><th>Cena</th></tr></thead>"
        "<tbody><tr><td>Ala</td><td>3,14</td></tr><tr><td>Kot</td><td>2,50</td></tr></tbody></table></section>"
    )
    assert "Ukryty" not in doc.html
    assert doc.warnings == ["sheet 'Ukryty' is hidden, skipped"]


def test_xlsm_same_bytes_same_html(tmp_path: Path) -> None:
    path_xlsx = tmp_path / "book.xlsx"
    path_xlsm = tmp_path / "book.xlsm"
    _make_xlsx(path_xlsx)
    path_xlsm.write_bytes(path_xlsx.read_bytes())
    assert _compact(convert_xlsx(path_xlsm, OPTS).html) == _compact(
        convert_xlsx(path_xlsx, OPTS).html
    )


def test_xlsx_row_limit_raises_during_iteration(tmp_path: Path) -> None:
    path = tmp_path / "book.xlsx"
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["nagłówek"])
    for i in range(2):
        ws.append([f"w{i}"])
    wb.save(path)
    with pytest.raises(LimitExceeded) as excinfo:
        convert_xlsx(path, ConvertOptions(max_sheet_rows=2))
    assert "sheet rows 3 > 2" in str(excinfo.value)


def test_xlsx_empty_workbook_raises(tmp_path: Path) -> None:
    path = tmp_path / "book.xlsx"
    openpyxl.Workbook().save(path)
    with pytest.raises(EmptyDocument) as excinfo:
        convert_xlsx(path, OPTS)
    assert "spreadsheet has no data" in str(excinfo.value)


def test_xlsx_cfb_container_encrypted(tmp_path: Path) -> None:
    path = tmp_path / "book.xlsx"
    path.write_bytes(_CFB + b"\x00" * 16)
    with pytest.raises(EncryptedFile):
        convert_xlsx(path, OPTS)


# ---------------------------------------------------------------- xls


def test_xls_fixture_sections_and_dates() -> None:
    doc = convert_xls(FIXTURE_XLS, OPTS)
    assert _compact(doc.html) == (
        '<section data-sheet="Dane"><table><thead><tr>'
        "<th>Nazwa</th><th>Kwota</th><th>Data</th></tr></thead>"
        "<tbody><tr><td>Ala</td><td>3.5</td><td>2024-01-02</td></tr>"
        "<tr><td>Kot</td><td>4</td><td>2024-02-03</td></tr></tbody></table></section>"
    )
    assert "Ukryty" not in doc.html
    assert doc.warnings == ["sheet 'Ukryty' is hidden, skipped"]


def test_xls_garbage_raises_corrupted(tmp_path: Path) -> None:
    path = tmp_path / "smieci.xls"
    path.write_bytes(b"to nie jest arkusz xls, tylko smieci" * 4)
    with pytest.raises(CorruptedFile):
        convert_xls(path, OPTS)


def test_xls_encrypted_message(monkeypatch: pytest.MonkeyPatch) -> None:
    def _boom(*args: object, **kwargs: object) -> None:
        raise xlrd.XLRDError("Workbook is encrypted")

    monkeypatch.setattr(xlrd, "open_workbook", _boom)
    with pytest.raises(EncryptedFile):
        convert_xls(Path("cokolwiek.xls"), OPTS)


# ---------------------------------------------------------------- csv/tsv


def test_csv_cp1250_semicolon_sniffed(tmp_path: Path) -> None:
    path = tmp_path / "d.csv"
    data = "Nazwa;Kwota\nAla;3,14\nKot;2,50\nZażółć;0,00\n"
    path.write_bytes(data.encode("cp1250"))
    doc = convert_csv(path, OPTS)
    assert doc.metadata["encoding"] == "cp1250"
    assert _compact(doc.html) == (
        "<table><thead><tr><th>Nazwa</th><th>Kwota</th></tr></thead><tbody>"
        "<tr><td>Ala</td><td>3,14</td></tr><tr><td>Kot</td><td>2,50</td></tr>"
        "<tr><td>Zażółć</td><td>0,00</td></tr></tbody></table>"
    )


def test_csv_utf8_bom_comma(tmp_path: Path) -> None:
    path = tmp_path / "d.csv"
    path.write_bytes(b"\xef\xbb\xbfa,b\n1,2\n")
    doc = convert_csv(path, OPTS)
    assert doc.metadata["encoding"] == "utf-8-sig"
    assert _compact(doc.html) == (
        "<table><thead><tr><th>a</th><th>b</th></tr></thead>"
        "<tbody><tr><td>1</td><td>2</td></tr></tbody></table>"
    )


def test_csv_quoted_field_with_newline_kept(tmp_path: Path) -> None:
    path = tmp_path / "d.csv"
    path.write_bytes(b'a,b\n1,"x\ny"\n')
    doc = convert_csv(path, OPTS)
    # esc() normalizuje tekst, ale \n zostaje — komórka zawiera dosłowny nowy wiersz
    assert "<td>x\ny</td>" in doc.html
    assert "<th>a</th>" in doc.html


def test_csv_sniffer_fallback_single_column(tmp_path: Path) -> None:
    path = tmp_path / "d.csv"
    path.write_bytes(b"a\nb\nc\n")
    doc = convert_csv(path, OPTS)
    assert _compact(doc.html) == (
        "<table><thead><tr><th>a</th></tr></thead>"
        "<tbody><tr><td>b</td></tr><tr><td>c</td></tr></tbody></table>"
    )


def test_csv_sniffer_fallback_semicolon_preferred(tmp_path: Path) -> None:
    path = tmp_path / "d.csv"
    # Sniffer nie ustala separatora (różne układy w wierszach); pierwsza linia ma więcej ';' niż ','
    path.write_bytes(b"a;b\nc\nd\n")
    doc = convert_csv(path, OPTS)
    assert _compact(doc.html) == (
        "<table><thead><tr><th>a</th><th>b</th></tr></thead>"
        "<tbody><tr><td>c</td><td></td></tr><tr><td>d</td><td></td></tr></tbody></table>"
    )


def test_tsv_two_columns(tmp_path: Path) -> None:
    path = tmp_path / "d.tsv"
    path.write_bytes(b"a\tb\n1\t2\n")
    doc = convert_tsv(path, OPTS)
    assert _compact(doc.html) == (
        "<table><thead><tr><th>a</th><th>b</th></tr></thead>"
        "<tbody><tr><td>1</td><td>2</td></tr></tbody></table>"
    )


def test_tsv_row_limit_raises(tmp_path: Path) -> None:
    path = tmp_path / "d.tsv"
    path.write_bytes(b"a\tb\n1\t2\n3\t4\n")
    with pytest.raises(LimitExceeded):
        convert_tsv(path, ConvertOptions(max_sheet_rows=1))


def test_csv_zero_bytes_raises(tmp_path: Path) -> None:
    path = tmp_path / "d.csv"
    path.write_bytes(b"")
    with pytest.raises(EmptyDocument):
        convert_csv(path, OPTS)


# ---------------------------------------------------------------- pickle


@pytest.mark.parametrize(
    "fn",
    [convert_xlsx, convert_xls, convert_csv, convert_tsv],
    ids=["xlsx", "xls", "csv", "tsv"],
)
def test_converts_are_picklable(fn: object) -> None:
    assert pickle.loads(pickle.dumps(fn)) is fn  # noqa: S301 — round-trip własnej funkcji
